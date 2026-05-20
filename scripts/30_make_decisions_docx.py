"""Generate a Word doc summarizing the multi-seed walk-back, the full method
search results, and the two decisions for the co-author."""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "BMVC_decisions.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def H1(t): return doc.add_heading(t, level=1)
def H2(t): return doc.add_heading(t, level=2)
def H3(t): return doc.add_heading(t, level=3)


def P(text=""):
    return doc.add_paragraph(text)


def P_runs(*runs):
    p = doc.add_paragraph()
    for text, kw in runs:
        r = p.add_run(text)
        r.bold = kw.get("bold", False)
        r.italic = kw.get("italic", False)
    return p


# Title
title = doc.add_heading("BMVC paper — multi-seed walk-back & two decisions for you", level=0)
sub = doc.add_paragraph()
sub.add_run("Date: 2026-05-08 · ").italic = True
sub.add_run("Re: follow-up to the IDEP doc I sent yesterday").italic = True

# Section 1 — what changed
H2("What changed since the IDEP doc")
P("I ran multi-seed validation on the SIEP method (5 seeds × 3 configs). The headline κ = 0.184 number from the doc was the lucky-seed result — multi-seed mean drops to 0.115 ± 0.047, within noise of the linear baseline. So the original method-paper claim (\"SIEP breaks the ceiling\") doesn't survive proper validation.")
P()
P_runs(
    ("I then built SIEP v2 with stabilization (cosine LR, EMA, longer training) and two contrastive formulations (SupCon and subject-stratified InfoNCE). The best honest result: ", {}),
    ("subject-stratified contrastive with cw=2.0 on DINOv2 gives mean κ_q = 0.138 ± 0.038 across 3 seeds.", {"bold": True}),
    (" All three seeds exceeded the linear baseline (0.107). It's the most stable variant we've found, but the gain is small.", {}),
)

# Section 2 — full picture
H2("Full picture across all 8 method categories")
P("All numbers on full 1,784-clip DAiSEE test, subject-disjoint. Best result per family:")

table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
for i, h in enumerate(["Method family", "Best κ_q"]):
    hdr[i].text = h
    for run in hdr[i].paragraphs[0].runs:
        run.bold = True

rows = [
    ("Zero-shot prompts (4 VLMs × 3 prompts)", "0.03"),
    ("Linear probes (LR / Ridge ordinal)", "0.107"),
    ("Multi-frame averaging", "0.103"),
    ("Calibration (CBU)", "0.104"),
    ("Linear concept erasure (LEACE)", "0.04 (worse)"),
    ("IDEP linear residualization", "0.124 (single-seed only)"),
    ("Non-linear adversarial (SIEP DANN)", "0.115 ± 0.047 (multi-seed)"),
    ("Non-linear contrastive (SIEP-SC) — best", "0.138 ± 0.038 (multi-seed)"),
]
for r in rows:
    cells = table.add_row().cells
    for i, v in enumerate(r):
        cells[i].text = v
    if "best" in r[0].lower():
        for cell in cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True

P()
P_runs(
    ("The κ ≈ 0.10–0.14 ceiling holds across ", {}),
    ("every", {"italic": True}),
    (" method family. The 0.107 → 0.138 lift from contrastive is the single positive movement in the matrix.", {}),
)

# Decision 1
H2("Decision 1 — Paper framing")

H3("Option A. \"The Engagement Ceiling\" — negative-result with exhaustive method coverage")
P("Eight method families tested rigorously, all bounded by κ ≈ 0.10–0.14. Centerpiece is the identity-engagement entanglement curve (asymmetric: identity high-rank redundant, engagement low-rank fragile). SIEP-contrastive becomes the strongest baseline, not the headline. Mechanistic explanation, not measurement. Title leans paradigm-critique.")

H3("Option B. \"SIEP-Contrastive\" — method paper with honest modest gain")
P("Propose SIEP-contrastive as the contribution. Headline is κ = 0.138 ± 0.038. Claims a small but stable improvement over linear methods, supported by the entanglement diagnosis. Smaller framing; reviewers can attack the magnitude.")

P()
P_runs(
    ("My read: ", {}),
    ("A is stronger.", {"bold": True}),
    (" The 0.031 absolute gain from SIEP-contrastive isn't large enough to anchor a method paper, and the ", {}),
    ("exhaustive search + geometric explanation", {"italic": True}),
    (" is genuinely novel. Reviewers can pick at \"modest method gain\"; they cannot pick at \"we tested every category of frozen-feature method, here's why none escapes.\" Want your honest read.", {}),
)

# Decision 2
H2("Decision 2 — Cross-dataset replication (FER2013 / RAF-DB)")
P("Whether to spend ~3–4 days adding FER2013 + RAF-DB cross-dataset experiments. This tests if the entanglement is engagement-specific (sharper claim about engagement labels) or general (broader claim about how VLM encoders allocate capacity for any subject-correlated affect task). Either outcome is publishable and pre-empts the strongest reviewer attack: \"this is a DAiSEE-specific quirk.\"")
P()
P("Cost: 3–4 days of the 21-day runway. Yes if you think the cross-dataset claim is worth slowing the writing; no if you'd rather invest those days in figures/writing/ethics/template.")
P()
P_runs(
    ("My lean: ", {}),
    ("yes, do FER2013 at minimum.", {"bold": True}),
    (" The entanglement claim doubles in strength if it generalizes (or specifies sharply if it doesn't). 3 days is affordable.", {}),
)

# Open ask
H2("Open ask")
P_runs(
    ("Call A or B on framing, yes or no on FER2013. Once you respond I lock the direction and start the BMVC template + intro draft.", {}),
)
P_runs(
    ("LLaVA + GPT-4o reruns, IDEP curve figure, SIEP variants, and the master results table are all done — paper materials are ready when framing is locked.", {"italic": True}),
)

doc.save(OUT)
print(f"Wrote {OUT}")
