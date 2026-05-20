"""Brief progress doc: 1 short summary + 1 table (completed vs in-progress)."""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "BMVC_progress_brief_2026-05-16.docx")

d = Document()

t = d.add_heading("BMVC 2026 — Progress Brief (2026-05-16)", 0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER

d.add_paragraph(
    "Headline: 90+ method variants run over the past week on DAiSEE engagement. "
    "Frozen-feature ceiling holds at κ_q ≈ 0.23–0.24. Paper draft underway as a "
    "diagnostic + DREAM (negative) + one robust adaptation (per-clip + per-subject "
    "blend, +0.014 on average). Abstract due 2026-05-22; paper 2026-05-29."
)

tbl = d.add_table(rows=1, cols=2)
tbl.style = "Light Grid Accent 1"
hdr = tbl.rows[0].cells
hdr[0].text = "Status"
hdr[1].text = "Item"

rows = [
    ("Completed",
     "Multi-encoder probes (CLIP-B/32, CLIP-L/14, DINOv2, SigLIP-L, SigLIP-SO400M, TIPSv2, emotion-ViT)"),
    ("Completed", "Headline frozen-feature SOTA: bagged SigLIP-L LR + ordinal threshold → κ_q = 0.238 [0.191, 0.284]"),
    ("Completed", "Full 1,784-clip test rerun (A1) — addresses N=300 sampling-bias reviewer attack"),
    ("Completed", "Class-prior calibration suite (CBU, Menon, SLD-EM) — addresses class-imbalance attack"),
    ("Completed", "IDEP entanglement curve (k = 0…69) on CLIP / DINOv2 / SigLIP-L"),
    ("Completed", "Subject-ID probe — confirms >99% recoverability across encoders"),
    ("Completed", "FER2013 cross-task positive control — same encoders reach κ ≈ 0.55"),
    ("Completed", "Within-subject ranking analysis — per-subject Spearman ≈ 0.05 (load-bearing diagnostic)"),
    ("Completed", "DREAM (anchor-modulation personalization): best κ_q = 0.222 — constructive negative"),
    ("Completed", "Per-subject baseline assignment — single-bag κ_q = 0.242 (lucky); 5-bag mean 0.226 ± 0.012"),
    ("Completed", "Per-clip + per-subject E[y] blend, α=0.5 — best single bag κ_q = 0.250; 5-bag mean 0.228 ± 0.018 (+0.014 over clip-only)"),
    ("Completed", "Bagging variants (clip-bootstrap, RSB, subject-bootstrap, super-bag K=50×10, PCA, TTA)"),
    ("Completed", "Non-linear probes (MLP focal, CORN ordinal, adapter, MC dropout, mixup, soft-kappa loss)"),
    ("Completed", "Fusion / stacking (heterogeneous, rank/geom/power, Dirichlet, val-stacking, OOF-stacking)"),
    ("Completed", "Adaptation methods (LEACE, IDEP residualization, SIEP adversarial, subject-centered, NN-anchor, hierarchical)"),
    ("Completed", "Threshold-tuning lift (+0.025 over raw bag) verified by K-fold CV"),
    ("Completed", "Comprehensive negative-results catalogue (34 methods tested and ruled out)"),
    ("Completed", "Paper draft scaffolding (paper_bmvc/main.tex, 8 sections, figures)"),
    ("In progress", "DREAM probe variant runs (PID 98239 on script 131_dream_probe)"),
    ("In progress", "VideoMAE-base encoding (69% complete, 5,883/8,571 clips; single-LR on partial features gives κ ≈ 0.016 — Plan A appears negative)"),
    ("In progress", "VideoMAE bagged probe + per-subject (PID 12551, still running)"),
    ("Next", "SigLIP-L last-block fine-tune (script 152, READY) — Plan B for positive-method evidence"),
    ("Next", "Within-subject ρ histogram figure with real per-subject data"),
    ("Next", "Replace placeholder BMVC LaTeX template with the official one"),
    ("Next", "Multi-frame SigLIP-L headline number (still pending)"),
    ("Next", "Reviewer-attack rerun completion (A2–A8: GPT-4o refusal annotations, CavT scale fix, LLaVA precision-match)"),
    ("Next", "Co-author review of paper draft and decision on framing tightness"),
]
for s, item in rows:
    r = tbl.add_row().cells
    r[0].text = s
    r[1].text = item

d.save(OUT)
print(f"Saved: {OUT}")
