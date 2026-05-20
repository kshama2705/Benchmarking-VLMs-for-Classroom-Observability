"""Day 1 report Word doc for the co-author. Pulls actual results from JSON
files; fills 'pending' if missing."""

import os, json
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "BMVC_day1_report.docx")
PPEP_RESULTS = os.path.join(BASE, "results", "ppep", "ppep_results.json")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def H1(t): return doc.add_heading(t, level=1)
def H2(t): return doc.add_heading(t, level=2)
def H3(t): return doc.add_heading(t, level=3)
def P(t=""): return doc.add_paragraph(t)


def P_runs(*runs):
    p = doc.add_paragraph()
    for text, kw in runs:
        r = p.add_run(text)
        r.bold = kw.get("bold", False)
        r.italic = kw.get("italic", False)
    return p


# Title
doc.add_heading("BMVC paper — Day 1 progress report", level=0)
sub = doc.add_paragraph()
sub.add_run("Date: 2026-05-10 · Deadlines: Abstract 2026-05-22 (12 days), Paper 2026-05-29 (19 days)").italic = True

# Section 1: Literature dive
H2("1. Literature dive")
P_runs(
    ("Target paper: ", {}),
    ("TIPS (Maninis et al., ICLR 2025) / TIPSv2 (Cao et al., CVPR 2026).", {"bold": True}),
    (" arxiv: 2410.16512 (v1), 2604.12012 (v2). HuggingFace: google/tipsv2 collection.", {}),
)
P("Core method: contrastive image-text pretraining augmented with masked image modeling, producing spatially-aware patch-level features (32×32 patch grid for TIPSv2-B14). This is distinct from standard CLIP which only has a global CLS token plus low-resolution patch tokens.")

P_runs(
    ("Why TIPS connects to our problem: ", {"bold": True}),
    ("our entanglement diagnosis (κ ≈ 0.10 ceiling) suggests engagement signal is being mixed with identity in global CLS features. If engagement signal lives mostly in the face region while identity leaks through clothing/posture/background, then face-region patch pooling should disentangle them more cleanly than global pooling. TIPS's spatially-aware features are designed exactly for this kind of dense readout.", {}),
)

# Section 2: Day-1 experiments
H2("2. Day 1 experiments")

H3("2.1 Face bbox detection on all 8,571 DAiSEE frames")
P("OpenCV Haar cascade frontal-face detector with histogram-equalized greyscale input. Picked the largest face per frame; saved normalized bbox + image size per clip.")
P("Result: detection rate 8,344 / 8,571 = 97.4%. Remaining 2.6% (227 frames) fall back to mean-patch pooling (no face region available).")

H3("2.2 Patch-Pooled Engagement Probe (PPEP)")
P("Method:")
P("• Encode each DAiSEE frame with DINOv2 ViT-B/14 to get 256 patch tokens (16×16 grid, 768-dim each) plus the CLS token.")
P("• Project the face bbox onto the patch grid; classify each patch as 'face' if ≥30% of its area lies inside the bbox.")
P("• Compute four pooled embeddings per clip: face_feat (face patches only), bg_feat (background patches), mean_feat (all patches mean), cls_feat (the CLS token, our existing baseline).")
P("• Train subject-disjoint LogReg-balanced and Ridge-ordinal probes on each pooling. Report κ_q with bootstrap CIs.")

H3("PPEP results (DINOv2 ViT-B/14)")
PPEP_CONCAT = os.path.join(BASE, "results", "ppep", "ppep_concat_results.json")
if os.path.exists(PPEP_RESULTS):
    with open(PPEP_RESULTS) as f:
        d = json.load(f)

    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["Pooling", "Probe", "κ_q", "95% CI"]):
        table.rows[0].cells[i].text = h
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for pool_label in ["face_feat", "bg_feat", "cls_feat", "mean_feat"]:
        if pool_label not in d:
            continue
        for probe_kind, probe_name in [("logreg", "LR"), ("ridge_ordinal", "Ridge")]:
            m = d[pool_label][probe_kind]
            row = table.add_row().cells
            row[0].text = pool_label
            row[1].text = probe_name
            row[2].text = f"{m['kappa_q']:.3f}"
            row[3].text = f"[{m['kappa_q_ci'][0]:.3f}, {m['kappa_q_ci'][1]:.3f}]"

    P()
    P_runs(
        ("Conclusion: simple patch pooling underperforms the CLS token. ", {"bold": True}),
        ("The CLS attention-based integration is doing real work; mean-pooling over face or background patches discards too much. The face-localization hypothesis is rejected on DINOv2 ViT-B/14.", {}),
    )

    if os.path.exists(PPEP_CONCAT):
        H3("Concatenation experiments (DINOv2 ViT-B/14)")
        with open(PPEP_CONCAT) as f:
            cd = json.load(f)
        P("Testing whether patches carry unique info beyond CLS, even if pooling is naive.")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(["Config", "Probe", "κ_q", "95% CI"]):
            table.rows[0].cells[i].text = h
            for run in table.rows[0].cells[i].paragraphs[0].runs:
                run.bold = True
        ordering = ["cls_only", "face_only", "bg_only", "face+bg",
                    "face+cls", "bg+cls", "face+bg+cls", "face+bg+cls+mean"]
        for cfg in ordering:
            if cfg not in cd:
                continue
            for probe_kind, probe_name in [("logreg", "LR"), ("ridge_ordinal", "Ridge")]:
                m = cd[cfg][probe_kind]
                row = table.add_row().cells
                row[0].text = cfg
                row[1].text = probe_name
                row[2].text = f"{m['kappa_q']:.3f}"
                row[3].text = f"[{m['kappa_q_ci'][0]:.3f}, {m['kappa_q_ci'][1]:.3f}]"
        P()
        P_runs(
            ("Bright spot: ", {"bold": True}),
            ("face+cls concatenation with Ridge ordinal gives κ_q = 0.116 [0.073, 0.164] vs CLS-Ridge baseline 0.090 — about +29% relative. But LR is worse (0.085 vs 0.107), so the gain is not consistent across probe families. CIs overlap baseline. Single-seed; not yet a clean method finding.", {}),
        )
else:
    P("(PPEP encoding still running — results will be inserted on completion.)")

# Section 3: Day-1 implications + Day-2 plan
H2("3. Implication")
P("The result on DINOv2 was the 'pooling location doesn't matter' outcome — neither face nor background patches alone match the CLS token. Engagement and identity appear roughly uniformly mixed across the 16×16 patch grid at DINOv2's resolution. Naive concatenation doesn't fix it either (mixed across probe families).")
P_runs(
    ("Important nuance: ", {"bold": True}),
    ("face+cls Ridge gain (0.116 vs 0.090) suggests patches DO carry some unique information, just too noisy to extract cleanly with mean-pooling. This motivates trying ", {}),
    ("(a) finer-grained patches (TIPSv2 32×32)", {"italic": True}),
    (", ", {}),
    ("(b) attention-pooling instead of mean", {"italic": True}),
    (", or ", {}),
    ("(c) larger encoders with denser features (SigLIP-L)", {"italic": True}),
    (".", {}),
)

H2("4. Day 2 plan (revised)")
P("4.1 PPEP on TIPSv2-B14 (HuggingFace google/tipsv2-b14, 32×32 patch grid at 448 input — 4× finer spatial resolution than DINOv2). ~2-3 hours encoding + 5 min probe. Top priority.")
P("4.2 PPEP on SigLIP-L/16-256 — its CLS already gives κ=0.199 (best result we have); patches might further improve. ~2-3 hours.")
P("4.3 If PPEP on TIPS/SigLIP-L doesn't help, pivot to learnable attention pooling over patches (train a small attention head, subject-disjoint). ~2 hours.")
P("4.4 If still no win, document the negative result and add to the exhaustive method search section of the paper. The exploration cost is small (~6 hours overnight) and the negative result is still informative.")

H2("5. Risks / decisions still open with co-author")
P("• Framing call (A vs B) still pending; PPEP results may push toward B if face-pooling gives a clean win.")
P("• OpenAI API key rotation pending (sent in earlier transcript).")
P("• FER2013 cross-dataset experiments already done overnight, results in BMVC_decisions doc.")

doc.save(OUT)
print(f"Wrote {OUT}")
